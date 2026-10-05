# Собирает базу знаний в двух видах: Excel (для заказчика) и Markdown (для бота).
from kb_data import *
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

def price(p): return f"{p:,}".replace(",", " ") + " ₽"
def num(v): return "—" if v is None else str(v).replace(".", ",")

sheets = {
    "О компании": (["Раздел", "Информация"], COMPANY),
    "Каталог": (["Категория", "Модель", "Мощность, кВт", "Производительность, м³/мин", "Давление, бар", "Цена", "Наличие"],
                [(c, m, num(k), num(q), num(b), price(p), s) for c, m, k, q, b, p, s in CATALOG] + [("Примечание", CATALOG_NOTE, "", "", "", "", "")]),
    "Услуги": (["Услуга", "Цена", "Описание"], SERVICES),
    "FAQ": (["Вопрос", "Ответ"], FAQ),
    "Квалификация": (["Что выяснить", "Как спросить / что нужно"], QUALIFY),
    "Скрипты": (["Ситуация", "Текст"], SCRIPTS),
    "Tone of Voice": (["Правило", "Описание"], TONE),
    "Запреты": (["Бот НЕ должен"], [(f,) for f in FORBIDDEN]),
}
wb = Workbook(); wb.remove(wb.active)
head = PatternFill("solid", fgColor="1F4E79")
for name, (cols, rows) in sheets.items():
    ws = wb.create_sheet(name)
    ws.append(cols)
    for r in rows: ws.append(list(r))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = head
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(12, w + 2), 70)
        for c in col: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
wb.save("baza_znaniy_PromVozduh.xlsx")

L = ["# База знаний: ПромВоздух (демо-данные для учебного MVP)", ""]
L += ["## О компании"] + [f"- {k}: {v}" for k, v in COMPANY] + [""]
L += ["## Каталог", "Формат: модель — мощность кВт / производительность м³/мин / давление бар — цена — наличие"]
for c, m, k, q, b, p, s in CATALOG:
    spec = " / ".join(x for x in [k and f"{num(k)} кВт", q and f"{num(q)} м³/мин", f"{num(b)} бар"] if x)
    L.append(f"- {c} {m} — {spec} — {price(p)} — {s}")
L += [CATALOG_NOTE, "", "## Услуги"] + [f"- {a} — {b}. {c}" for a, b, c in SERVICES] + [""]
L += ["## Частые вопросы"] + [f"- В: {q}\n  О: {a}" for q, a in FAQ] + [""]
L += ["## Что выяснить для заявки"] + [f"- {a}: {b}" for a, b in QUALIFY] + [""]
L += ["## Скрипты"] + [f"- {a}: {b}" for a, b in SCRIPTS] + [""]
L += ["## Тон общения"] + [f"- {a}: {b}" for a, b in TONE] + [""]
L += ["## Запрещено"] + [f"- {f}" for f in FORBIDDEN]
open("knowledge_base.md", "w").write("\n".join(L) + "\n")
print("ok")
